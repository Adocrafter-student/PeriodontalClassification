"""Build editable DOCX and exportable methodology figures; no model training."""
from pathlib import Path
import re
import sys

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE / '.build_deps'))
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from matplotlib.patches import FancyBboxPatch, FancyArrowPatch
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

FIGURES = HERE / 'figures'
FIGURES.mkdir(exist_ok=True)
plt.rcParams.update({'font.family': 'DejaVu Sans', 'svg.fonttype': 'none'})
COLORS = {'data': '#edf3f8', 'cnn': '#e4eff9', 'dino': '#e4f2ed',
          'test': '#fff0d8', 'note': '#f1f1f1'}


def canvas(height):
    fig, ax = plt.subplots(figsize=(7.2, height))
    ax.set(xlim=(0, 7.2), ylim=(0, height))
    ax.axis('off')
    fig.subplots_adjust(left=0, right=1, bottom=0, top=1)
    return fig, ax


def box(ax, x, y, w, h, text, kind='data', size=10):
    ax.add_patch(FancyBboxPatch((x-w/2, y-h/2), w, h,
                 boxstyle='round,pad=0.025,rounding_size=0.07', linewidth=.9,
                 edgecolor='#415a6b', facecolor=COLORS[kind], zorder=2))
    ax.text(x, y, text, ha='center', va='center', fontsize=size,
            color='#152d3a', linespacing=1.4, zorder=3)


def arrow(ax, start, end, dashed=False):
    ax.add_patch(FancyArrowPatch(start, end, arrowstyle='-|>', mutation_scale=11,
                 linewidth=1, color='#415a6b', linestyle='--' if dashed else '-', zorder=1))


def export(fig, name):
    for suffix in ('png', 'svg', 'pdf'):
        fig.savefig(FIGURES / f'{name}.{suffix}', dpi=300, facecolor='white')
    plt.close(fig)


def build_figures():
    fig, ax = canvas(5.6)
    box(ax, 1.9, 4.95, 3.25, .85, 'Local BRAR archive\n988 images across levels 1, 2 and 3')
    box(ax, 5.4, 4.95, 2.9, .85, 'Professor reference file\n252 image IDs + stages I–IV')
    box(ax, 3.6, 3.65, 5.45, .85, 'Match numerical image IDs\nCheck metadata, image readability and duplicates')
    arrow(ax, (1.9, 4.5), (3.0, 4.1)); arrow(ax, (5.4, 4.5), (4.2, 4.1))
    box(ax, 2.0, 2.15, 3.0, 1.1, 'Analyzed: 252 labeled images\nI: 10   II: 53   III: 140   IV: 49\nAll matched original level_3', 'dino', 9.5)
    box(ax, 5.4, 2.15, 2.8, 1.1, 'Excluded: 736 images\nNo professor stage label\nNo feature extraction or fitting', 'note', 9.5)
    arrow(ax, (2.6, 3.2), (2.0, 2.75)); arrow(ax, (4.6, 3.2), (5.4, 2.75))
    box(ax, 2.0, .8, 3.0, .7, 'Five shared outer folds\nOne held-out prediction per image', 'test', 9.5)
    arrow(ax, (2.0, 1.55), (2.0, 1.2))
    ax.text(5.4, .8, 'Metadata: audit only\nModel inputs: images only', ha='center', va='center', fontsize=9, color='#415a6b')
    export(fig, 'figure_5_1_dataset')

    fig, ax = canvas(7.5)
    box(ax, 3.6, 7.0, 5.6, .65, 'One of five shared outer folds', 'data', 11)
    box(ax, 2.35, 5.95, 3.8, .65, 'Development set\n201 or 202 images', 'data')
    box(ax, 5.8, 5.95, 2.05, .65, 'Held-out test set\n51 or 50 images', 'test', 9.5)
    arrow(ax, (2.7, 6.65), (2.35, 6.3)); arrow(ax, (4.8, 6.65), (5.8, 6.3))
    box(ax, 1.7, 4.55, 2.8, 1.2, 'EfficientNet selection\n170/171 fitting images\n31 validation images\nSelect epoch by macro F1', 'cnn', 9.5)
    box(ax, 4.75, 4.55, 2.8, 1.2, 'DINOv2 selection\nThree inner folds\nFit scaler within each fold\nSelect C by macro F1', 'dino', 9.5)
    arrow(ax, (1.7, 5.6), (1.7, 5.2)); arrow(ax, (3.1, 5.6), (4.75, 5.2))
    box(ax, 1.7, 2.75, 2.8, 1.05, 'Reset external initialization\nRefit encoder and head on\nall development images\nfor selected epoch count', 'cnn', 9.3)
    box(ax, 4.75, 2.75, 2.8, 1.05, 'Encoder remains frozen\nRefit scaler and classifier on\nall development images\nusing selected C', 'dino', 9.3)
    arrow(ax, (1.7, 3.9), (1.7, 3.32)); arrow(ax, (4.75, 3.9), (4.75, 3.32))
    box(ax, 3.6, 1.2, 5.6, .8, 'Predict the held-out fold once\nPool predictions after repeating all five folds', 'test', 10)
    arrow(ax, (1.7, 2.18), (2.4, 1.64)); arrow(ax, (4.75, 2.18), (4.65, 1.64))
    arrow(ax, (6.9, 5.95), (6.9, 1.2), True); arrow(ax, (6.9, 1.2), (6.45, 1.2), True)
    ax.text(3.6, .28, 'Outer test labels are never used to select C or the training duration.',
            ha='center', va='center', fontsize=9, color='#415a6b')
    export(fig, 'figure_5_2_validation')

    fig, ax = canvas(6.9)
    box(ax, 3.6, 6.4, 5.8, .7, 'Panoramic image → aspect-ratio-preserving input\n448 × 448 pixels, checkpoint-specific normalization', 'data', 10)
    xs = [1.25, 3.6, 5.95]
    labels = ['EfficientNet-B4\nWhole image\nImageNet initialization',
              'DINOv2 ViT-S/14\nWhole image\nFrozen encoder',
              'DINOv2 ViT-S/14\nWhole + 3 regions\nFrozen encoder']
    for x, label, kind in zip(xs, labels, ['cnn', 'dino', 'dino']):
        box(ax, x, 5.0, 2.12, .95, label, kind, 9.5)
        arrow(ax, (x, 6.0), (x, 5.52))
    mids = ['Global pooled\nconvolutional features', 'Whole-image vector\n384 features',
            'Whole vector + mean\nof regional vectors\n768 features']
    for x, label, kind in zip(xs, mids, ['cnn', 'dino', 'dino']):
        box(ax, x, 3.5, 2.12, .85, label, kind, 9.5)
        arrow(ax, (x, 4.48), (x, 3.97))
    bottoms = ['Train encoder +\n256-unit head\nWeighted cross entropy',
               'Fit standardization +\nweighted logistic\nregression',
               'Fit standardization +\nweighted logistic\nregression']
    for x, label, kind in zip(xs, bottoms, ['cnn', 'dino', 'dino']):
        box(ax, x, 2.0, 2.12, 1.0, label, kind, 9.5)
        arrow(ax, (x, 3.03), (x, 2.55))
        arrow(ax, (x, 1.45), (x, .95))
    box(ax, 3.6, .6, 6.7, .62, 'Four stage probabilities → most probable reference stage\nEvaluation uses held-out predictions, not predictions on fitting images', 'test', 9.2)
    export(fig, 'figure_5_3_models')


def inline(paragraph, text):
    pattern = r'(\[[^\]]+\]\(https?://[^)]+\)|\*\*[^*]+\*\*)'
    for part in re.split(pattern, text):
        match = re.fullmatch(r'\[([^\]]+)\]\((https?://[^)]+)\)', part)
        if match:
            link = OxmlElement('w:hyperlink')
            link.set(qn('r:id'), paragraph.part.relate_to(match[2], RT.HYPERLINK, is_external=True))
            run = OxmlElement('w:r'); props = OxmlElement('w:rPr')
            color = OxmlElement('w:color'); color.set(qn('w:val'), '235789'); props.append(color)
            run.append(props); t = OxmlElement('w:t'); t.text = match[1]; run.append(t)
            link.append(run); paragraph._p.append(link)
        elif part.startswith('**') and part.endswith('**'):
            paragraph.add_run(part[2:-2]).bold = True
        else:
            paragraph.add_run(part)


def document():
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(8.27), Inches(11.69)
    sec.top_margin = sec.bottom_margin = Inches(.8)
    sec.left_margin = sec.right_margin = Inches(.85)
    normal = doc.styles['Normal']
    normal.font.name = 'Times New Roman'; normal.font.size = Pt(11)
    normal.paragraph_format.line_spacing = 1.15
    normal.paragraph_format.space_after = Pt(6)
    for level in (1, 2):
        style = doc.styles[f'Heading {level}']
        style.font.name = 'Calibri'
        style.font.size = Pt(16 if level == 1 else 12)
        style.font.color.rgb = RGBColor.from_string('193D52')
        style.paragraph_format.keep_with_next = True
    header = sec.header.paragraphs[0]
    header.text = 'BRAR stage prediction | Thesis section draft'
    header.style = doc.styles['Caption']
    footer = sec.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    field = OxmlElement('w:fldSimple'); field.set(qn('w:instr'), 'PAGE'); footer._p.append(field)
    doc.add_heading('Prediction of Professor-Assigned Periodontal Stages from BRAR Panoramic Radiographs', 0)
    doc.add_paragraph('EfficientNet-B4 with frozen DINOv2 comparisons', style='Subtitle')
    doc.add_paragraph('Draft of the missing thesis sections, based on the existing literature review and verified five-fold results.')
    doc.add_paragraph('Chapter 2 remains the supplied literature review and is not reproduced here. Integration notes are supplied separately. This draft contains Introduction, Research Questions / Hypotheses, Dataset, Methodology, Results, Discussion, Conclusion, and references for these sections.')
    doc.add_page_break()
    lines = (HERE / 'missing_sections.md').read_text(encoding='utf-8').splitlines()
    index = 0
    while index < len(lines):
        line = lines[index]
        if not line.strip():
            index += 1; continue
        if line.startswith('# '):
            p = doc.add_heading(line[2:], level=1)
            if not line.startswith('# 1.'):
                p.paragraph_format.page_break_before = True
        elif line.startswith('## '):
            doc.add_heading(line[3:], level=2)
        elif line.startswith('!['):
            source = re.search(r'\]\(([^)]+)\)', line)[1]
            p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.keep_with_next = True
            p.add_run().add_picture(str(HERE / source), width=Inches(6.45))
        elif line.startswith('|'):
            rows = []
            while index < len(lines) and lines[index].startswith('|'):
                row = [v.strip() for v in lines[index].strip('|').split('|')]
                if not all(re.fullmatch(r':?-+:?', v) for v in row):
                    rows.append(row)
                index += 1
            table = doc.add_table(rows=1, cols=len(rows[0])); table.style = 'Light Shading Accent 1'
            repeat = OxmlElement('w:tblHeader'); table.rows[0]._tr.get_or_add_trPr().append(repeat)
            for i, row in enumerate(rows):
                cells = table.rows[0].cells if i == 0 else table.add_row().cells
                for cell, value in zip(cells, row):
                    cell.text = value
                    for p in cell.paragraphs:
                        for run in p.runs:
                            run.font.size = Pt(9)
                            if i == 0: run.bold = True
                no_split = OxmlElement('w:cantSplit'); cells[0]._tc.getparent().get_or_add_trPr().append(no_split)
            doc.add_paragraph()
            continue
        else:
            p = doc.add_paragraph()
            if line.startswith(('Figure ', 'Table ')):
                p.style = doc.styles['Caption']
                if line.startswith('Table '): p.paragraph_format.keep_with_next = True
            inline(p, line)
        index += 1
    path = HERE / 'missing_sections.docx'
    doc.save(path)
    print(f'Saved {path}')
    # Structural export checks: all figures embedded and all requested sections present.
    check = Document(path)
    expected_figures = sum(line.startswith('![') for line in lines)
    assert len(check.inline_shapes) == expected_figures
    assert len(check.tables) == 3
    headings = [p.text for p in check.paragraphs if p.style.name == 'Heading 1']
    assert len(headings) == 8, headings
    print(f'Verified {len(headings)} chapter/reference headings, {expected_figures} embedded figures, 3 tables.')


if __name__ == '__main__':
    build_figures()
    document()
